"""Reading and writing .xlsx workbooks (openpyxl)."""

import io
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal

from django.utils import translation
from django.utils.translation import gettext_lazy as _
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from common.exceptions import DomainError

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MAX_ROWS = 20_000


class BadWorkbook(DomainError):
    code = "BAD_WORKBOOK"
    default_detail = _("This file is not an Excel workbook (.xlsx) this system can read.")


@dataclass(frozen=True)
class Column:
    key: str
    label: str  # lazy translation
    required: bool = False
    width: int = 16
    kind: str = "text"  # text | date | datetime | money | int | bool
    example: object = ""


@dataclass
class Sheet:
    title: str
    columns: list[Column]
    rows: list[list] = field(default_factory=list)


def _names(column: Column) -> set[str]:
    """Every header that means this column: its key and its label in each language."""
    names = {column.key.lower(), column.key.replace("_", " ").lower()}
    for language in ("en", "ko-kp"):
        with translation.override(language):
            names.add(str(column.label).strip().lower())
    return names


def read_rows(upload, columns: list[Column]) -> list[tuple[int, dict]]:
    """The first sheet's rows as (Excel row number, {key: value}); the first row is the
    header (keys or labels, any order, in English or Korean). Blank rows are skipped;
    unknown columns are ignored."""
    try:
        book = load_workbook(upload, read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001  (zip, xml and format errors alike)
        raise BadWorkbook() from exc
    sheet = book.worksheets[0]
    rows = sheet.iter_rows(values_only=True)
    header = next(rows, None) or ()
    lookup = {name: column.key for column in columns for name in _names(column)}
    keys = [lookup.get(str(cell).strip().lower()) if cell is not None else None for cell in header]
    missing = [str(c.label) for c in columns if c.required and c.key not in keys]
    if missing:
        raise BadWorkbook(
            _("Missing columns: %(columns)s.") % {"columns": ", ".join(missing)},
            details={"missing": missing},
        )
    out = []
    for number, values in enumerate(rows, start=2):
        if number - 1 > MAX_ROWS:
            raise BadWorkbook(_("At most %(max)s rows per file.") % {"max": MAX_ROWS})
        row = {}
        for key, value in zip(keys, values, strict=False):
            if key is None:
                continue
            if isinstance(value, str):
                value = value.strip()
            row[key] = None if value == "" else value
        if any(v is not None for v in row.values()):
            out.append((number, row))
    book.close()
    return out


def _cell(value, kind: str):
    if value is None:
        return None
    if kind == "money":
        return Decimal(str(value))
    if kind == "datetime" and isinstance(value, datetime):
        return value.replace(tzinfo=None)
    return value


FORMATS = {"money": "#,##0.00", "date": "yyyy-mm-dd", "datetime": "yyyy-mm-dd hh:mm"}


def workbook(sheets: list[Sheet], text_rows: int = 0) -> bytes:
    """Sheets with a bold header row, frozen, filterable, with fitting column widths.
    `text_rows`: format that many rows of the text columns as Text (templates), so Excel
    keeps leading zeros in employee numbers and card UIDs."""
    book = Workbook()
    book.remove(book.active)
    for spec in sheets:
        ws = book.create_sheet(str(spec.title)[:31])
        ws.append([str(c.label) for c in spec.columns])
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="E8E8E8")
        for row in spec.rows:
            ws.append([_cell(v, c.kind) for v, c in zip(row, spec.columns, strict=True)])
        for index, column in enumerate(spec.columns, start=1):
            letter = get_column_letter(index)
            ws.column_dimensions[letter].width = column.width
            if column.kind in FORMATS:
                for (cell,) in ws.iter_rows(min_row=2, min_col=index, max_col=index):
                    cell.number_format = FORMATS[column.kind]
            elif column.kind == "text" and text_rows:
                for row in range(2, text_rows + 2):
                    ws.cell(row=row, column=index).number_format = "@"
        ws.freeze_panes = "A2"
        if spec.rows:
            ws.auto_filter.ref = ws.dimensions
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def template(title: str, columns: list[Column]) -> bytes:
    """An import template: the header and one example row."""
    return workbook([Sheet(title, columns, [[c.example for c in columns]])], text_rows=1000)


def as_date(value) -> date | str | None:
    """Excel dates come as datetime; text dates are left for the serializer to parse."""
    if isinstance(value, datetime):
        return value.date()
    return value
